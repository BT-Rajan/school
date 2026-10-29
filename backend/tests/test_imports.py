import csv
import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models import MailMessage, MailPriority, MfaMethod, Role, Teacher

from .helpers import make_cluster, make_school, seed_reference, signed_in

TEACHER_HEADER_EN = [
    "Teacher ID",
    "Teacher name",
    "Email",
    "School ID",
    "Designation",
    "Subjects",
    "Classes",
    "Experience (years)",
    "Qualifications",
]
TEACHER_HEADER_AR = [
    "رقم المعلم",
    "اسم المعلم",
    "البريد الإلكتروني",
    "رقم المدرسة",
    "المسمى الوظيفي",
    "المواد",
    "الصفوف",
    "سنوات الخبرة",
    "المؤهلات",
]


def _csv(header, rows) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


@pytest.fixture
def setup(new_client, outbox):
    seed_reference()
    client, head = signed_in(new_client, outbox, Role.CLUSTER_HEAD, "head@example.org")
    cluster = make_cluster("C1", head)
    make_school(cluster, "S1")
    other = make_cluster("C2")
    make_school(other, "S9")
    admin, _ = signed_in(new_client, outbox, Role.DISTRICT_ADMIN, "admin@example.org", mfa=MfaMethod.TOTP)
    return {"head": client, "admin": admin}


def _preview(client, entity, name, content):
    return client.post(f"/api/imports/{entity}/preview", files={"file": (name, content)})


def test_teacher_import_preview_commit_and_rerun(setup, db):
    rows = [
        ["T1", "Ahmed Ali", "ahmed@example.org", "S1", "Teacher", "Mathematics; Science", "Grade 1", "5", "B.Ed:2012"],
        # Arabic list values are accepted as well as English ones and codes.
        ["T2", "سارة محمد", "sara@example.org", "S1", "معلم", "اللغة العربية", "الصف الثاني", "3", "MED"],
        ["T3", "Bad Row", "bad@example.org", "S1", "Teacher", "Astrology", "Grade 1", "2", "B.Ed"],
        ["T4", "Other Cluster", "o@example.org", "S9", "Teacher", "Mathematics", "Grade 1", "2", "B.Ed"],
        ["T1", "Duplicate", "dup@example.org", "S1", "Teacher", "Mathematics", "Grade 1", "2", "B.Ed"],
    ]
    preview = _preview(setup["head"], "teachers", "teachers.csv", _csv(TEACHER_HEADER_EN, rows))
    assert preview.status_code == 200, preview.text
    result = preview.json()
    actions = [r["action"] for r in result["rows"]]
    assert actions == ["created", "created", "error", "error", "error"]
    assert result["rows"][2]["errors"][0]["code"] == "reference_unknown"
    assert result["rows"][3]["errors"][0]["code"] == "out_of_scope"
    assert result["rows"][4]["errors"][0]["code"] == "duplicate_in_file"
    # Preview changes nothing.
    assert db.execute(select(func.count()).select_from(Teacher)).scalar_one() == 0

    committed = setup["head"].post(f"/api/imports/{result['id']}/commit").json()
    assert committed["summary"] == {"created": 2, "error": 3}
    assert sum(day["count"] for day in committed["invitation_schedule"]) == 2
    db.rollback()  # end the test session's snapshot so it sees the commit
    assert db.execute(select(func.count()).select_from(Teacher)).scalar_one() == 2
    invitations = db.execute(
        select(func.count()).select_from(MailMessage).where(MailMessage.priority == MailPriority.INVITATION)
    ).scalar_one()
    assert invitations == 2
    # Committing twice is refused.
    assert setup["head"].post(f"/api/imports/{result['id']}/commit").status_code == 409

    # Re-running the same file is idempotent: nothing new is created.
    again = _preview(setup["head"], "teachers", "teachers.csv", _csv(TEACHER_HEADER_EN, rows[:2])).json()
    assert [r["action"] for r in again["rows"]] == ["unchanged", "unchanged"]

    # A changed value becomes an update.
    rows[0][5] = "Mathematics"
    changed = _preview(setup["head"], "teachers", "teachers.csv", _csv(TEACHER_HEADER_EN, rows[:1])).json()
    assert changed["rows"][0]["action"] == "updated"


def test_arabic_headings_and_xlsx(setup):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(TEACHER_HEADER_AR)
    sheet.append(["T7", "نورة", "noura@example.org", "S1", "T1", "MATH", "G1", 4, "BED"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = _preview(setup["head"], "teachers", "teachers.xlsx", buffer.getvalue()).json()
    assert result["rows"][0]["action"] == "created", result


def test_missing_and_unknown_columns(setup):
    r = _preview(setup["head"], "teachers", "t.csv", _csv(["Teacher ID", "Shoe size"], [["T1", "42"]]))
    assert r.status_code == 422
    detail = r.json()["error"]["detail"]
    assert "Teacher name" in detail["missing"] and detail["unknown"] == ["Shoe size"]


def test_non_utf8_csv_rejected(setup):
    content = "Teacher ID,Teacher name\nT1,Caf\xe9\n".encode("latin-1")
    r = _preview(setup["head"], "teachers", "t.csv", content)
    assert r.json()["error"]["code"] == "encoding_unsupported"


def test_school_import_idempotent(setup):
    header = ["Cluster ID", "School ID", "School name", "Address", "School type", "Medium of instruction", "Phone"]
    rows = [
        ["C1", "S2", "New School", "Road 1", "primary", "arabic", "+96512345678"],
        ["C1", "S1", "School S1", "Street 1", "primary", "arabic", "+96500000000"],
    ]
    result = _preview(setup["head"], "schools", "s.csv", _csv(header, rows)).json()
    assert [r["action"] for r in result["rows"]] == ["created", "unchanged"]
    setup["head"].post(f"/api/imports/{result['id']}/commit")
    again = _preview(setup["head"], "schools", "s.csv", _csv(header, rows)).json()
    assert [r["action"] for r in again["rows"]] == ["unchanged", "unchanged"]


def test_supervisor_import(setup):
    header = ["Staff ID", "Name", "Email", "Cluster ID"]
    result = _preview(
        setup["head"], "supervisors", "s.csv", _csv(header, [["SV1", "Sup One", "sv1@example.org", "C1"]])
    ).json()
    committed = setup["head"].post(f"/api/imports/{result['id']}/commit").json()
    assert committed["summary"] == {"created": 1}


def test_export_in_arabic_xlsx_with_district_code(setup):
    r = setup["admin"].get("/api/exports/schools", params={"format": "xlsx", "locale": "ar"})
    assert r.status_code == 200
    sheet = load_workbook(io.BytesIO(r.content)).active
    assert sheet.sheet_view.rightToLeft
    header = [c.value for c in sheet[1]]
    assert header[0] == "المنطقة" and "رقم المدرسة" in header
    codes = sorted(row[2] for row in sheet.iter_rows(min_row=2, values_only=True))
    assert codes == ["S1", "S9"]
    assert {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)} == {"KW-01"}


def test_cluster_head_export_is_scoped(setup):
    r = setup["head"].get("/api/exports/schools", params={"format": "csv"})
    text = r.content.decode("utf-8-sig")
    assert "S1" in text and "S9" not in text
