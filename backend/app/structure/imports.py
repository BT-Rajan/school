"""CSV and Excel import and export for schools, teachers and supervisors (FR-S4, FR-S5, FR-T2, FR-P2).

Import is two steps. Preview validates every row by running the real create
or update inside a savepoint that is rolled back, so preview and commit
apply exactly the same rules. Commit re-checks each row, because data may
have changed in between, and is idempotent on school ID or staff ID.
"""

import csv
import io
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import clock
from app.audit import service as audit
from app.auth.deps import Principal
from app.district import get_district
from app.errors import AppError, invalid, not_found
from app.mail import queue as mail
from app.models import Cluster, ImportBatch, Role, School, Supervisor, Teacher, User

from . import people, service
from .common import active_custom_fields, find_reference_by_label, fold_search, managed_cluster_ids
from .schemas import PersonUpdate, SchoolIn, SchoolUpdate, SupervisorIn, TeacherIn, TeacherUpdate

MAX_ROWS = 5000
ENTITIES = ("schools", "teachers", "supervisors")


@dataclass(frozen=True)
class Column:
    field: str
    en: str
    ar: str
    required: bool = True


COLUMNS: dict[str, list[Column]] = {
    "schools": [
        Column("cluster", "Cluster ID", "رقم المجمع"),
        Column("code", "School ID", "رقم المدرسة"),
        Column("name", "School name", "اسم المدرسة"),
        Column("address", "Address", "العنوان"),
        Column("school_type", "School type", "نوع المدرسة"),
        Column("medium", "Medium of instruction", "لغة التدريس"),
        Column("contact_phone", "Phone", "الهاتف"),
        Column("contact_email", "Email", "البريد الإلكتروني", required=False),
    ],
    "teachers": [
        Column("staff_id", "Teacher ID", "رقم المعلم"),
        Column("name", "Teacher name", "اسم المعلم"),
        Column("email", "Email", "البريد الإلكتروني"),
        Column("school", "School ID", "رقم المدرسة"),
        Column("designation", "Designation", "المسمى الوظيفي"),
        Column("subjects", "Subjects", "المواد"),
        Column("classes", "Classes", "الصفوف"),
        Column("experience_years", "Experience (years)", "سنوات الخبرة"),
        Column("qualifications", "Qualifications", "المؤهلات"),
        Column("locale", "Language", "اللغة", required=False),
    ],
    "supervisors": [
        Column("staff_id", "Staff ID", "الرقم الوظيفي"),
        Column("name", "Name", "الاسم"),
        Column("email", "Email", "البريد الإلكتروني"),
        Column("cluster", "Cluster ID", "رقم المجمع"),
        Column("locale", "Language", "اللغة", required=False),
    ],
}
CUSTOM_ENTITY = {"schools": "school", "teachers": "teacher"}


class RowError(Exception):
    def __init__(self, errors: list[dict[str, Any]]):
        self.errors = errors


# Parsing ----------------------------------------------------------------------------


def _read_table(filename: str, content: bytes) -> list[list[Any]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:  # noqa: BLE001 - any parse failure is a bad file
            raise invalid("file_unreadable") from exc
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    if name.endswith(".csv"):
        try:
            # FR-I8: UTF-8, with or without a byte order mark.
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise invalid("encoding_unsupported") from exc
        return [row for row in csv.reader(io.StringIO(text))]
    raise invalid("file_type_unsupported")


def _header_map(db: Session, entity: str, header: list[Any]) -> dict[int, str]:
    targets: dict[str, str] = {}
    for column in COLUMNS[entity]:
        for label in (column.field, column.en, column.ar):
            targets[fold_search(label)] = column.field
    if entity in CUSTOM_ENTITY:
        for field in active_custom_fields(db, CUSTOM_ENTITY[entity]):
            for label in (field.key, field.label_en, field.label_ar):
                targets[fold_search(label)] = f"custom.{field.key}"
    mapping: dict[int, str] = {}
    unknown = []
    for index, cell in enumerate(header):
        label = fold_search(str(cell or ""))
        if not label:
            continue
        if label in targets:
            mapping[index] = targets[label]
        else:
            unknown.append(str(cell))
    missing = [c.en for c in COLUMNS[entity] if c.required and c.field not in mapping.values()]
    if missing or unknown:
        raise invalid("columns_invalid", {"missing": missing, "unknown": unknown})
    return mapping


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def parse(db: Session, entity: str, filename: str, content: bytes) -> list[dict[str, Any]]:
    table = [row for row in _read_table(filename, content) if any(_cell(v) for v in row)]
    if not table:
        raise invalid("file_empty")
    if len(table) - 1 > MAX_ROWS:
        raise invalid("too_many_rows", {"max": MAX_ROWS})
    mapping = _header_map(db, entity, table[0])
    rows = []
    for number, raw in enumerate(table[1:], start=2):
        values: dict[str, Any] = {"custom": {}}
        for index, field in mapping.items():
            cell = _cell(raw[index]) if index < len(raw) else ""
            if field.startswith("custom."):
                if cell != "":
                    values["custom"][field.removeprefix("custom.")] = cell
            else:
                values[field] = cell
        rows.append({"row": number, "values": values})
    return rows


def _split(text: str) -> list[str]:
    for separator in ("،", "|", ","):
        text = text.replace(separator, ";")
    return [part.strip() for part in text.split(";") if part.strip()]


def _ref_codes(db: Session, list_name: str, text: str, field: str) -> list[str]:
    codes, unknown = [], []
    for label in _split(text):
        ref = find_reference_by_label(db, list_name, label)
        (codes.append(ref.code) if ref else unknown.append(label))
    if unknown:
        raise RowError([{"field": field, "code": "reference_unknown", "value": ", ".join(unknown)}])
    return codes


def _qualifications(db: Session, text: str) -> list[dict[str, Any]]:
    """Each entry is "name" or "name:year", for example "B.Ed:2012; M.Ed"."""
    result, unknown = [], []
    for part in _split(text):
        label, _, year = part.partition(":")
        ref = find_reference_by_label(db, "qualification", label)
        if ref is None:
            unknown.append(label)
            continue
        result.append({"code": ref.code, "year": int(year) if year.strip().isdigit() else None})
    if unknown:
        raise RowError([{"field": "qualifications", "code": "reference_unknown", "value": ", ".join(unknown)}])
    return result


# Row handling ---------------------------------------------------------------------------


def _cluster_by_code(db: Session, principal: Principal, code: str) -> Cluster:
    cluster = db.execute(select(Cluster).where(Cluster.code == code)).scalar_one_or_none()
    if cluster is None:
        raise RowError([{"field": "cluster", "code": "cluster_unknown", "value": code}])
    allowed = managed_cluster_ids(db, principal)
    if allowed is not None and cluster.id not in allowed:
        raise RowError([{"field": "cluster", "code": "out_of_scope", "value": code}])
    return cluster


def _school_by_code(db: Session, principal: Principal, code: str) -> School:
    school = db.execute(select(School).where(School.code == code)).scalar_one_or_none()
    if school is None:
        raise RowError([{"field": "school", "code": "school_unknown", "value": code}])
    allowed = managed_cluster_ids(db, principal)
    if allowed is not None and school.cluster_id not in allowed:
        raise RowError([{"field": "school", "code": "out_of_scope", "value": code}])
    return school


def _user_by_staff_id(db: Session, staff_id: str, role: Role) -> User | None:
    user = db.execute(select(User).where(User.staff_id == staff_id)).scalar_one_or_none()
    if user is not None and user.role != role:
        raise RowError([{"field": "staff_id", "code": "staff_id_taken"}])
    return user


def _apply_row(db: Session, principal: Principal, entity: str, values: dict[str, Any], send_at=None) -> str:
    """Create or update from one row. Returns created, updated or unchanged."""
    if entity == "schools":
        cluster = _cluster_by_code(db, principal, values.get("cluster", ""))
        fields = {
            k: values.get(k) or None
            for k in ("code", "name", "address", "school_type", "medium", "contact_phone", "contact_email")
        }
        existing = db.execute(select(School).where(School.code == fields["code"])).scalar_one_or_none()
        if existing is None:
            service.create_school(
                db, principal, SchoolIn(cluster_id=cluster.id, custom=values["custom"], **fields), commit=False
            )
            return "created"
        allowed = managed_cluster_ids(db, principal)
        if allowed is not None and existing.cluster_id not in allowed:
            raise RowError([{"field": "code", "code": "out_of_scope"}])
        if existing.cluster_id != cluster.id:
            raise RowError([{"field": "cluster", "code": "school_in_other_cluster"}])
        before = service.school_out(db, existing)
        update = {k: v for k, v in fields.items() if k != "code"}
        service.update_school(
            db,
            principal,
            existing.id,
            SchoolUpdate(version=existing.version, custom=values["custom"] or None, **update),
            commit=False,
        )
        return _changed(before, service.school_out(db, existing))

    if entity == "teachers":
        school = _school_by_code(db, principal, values.get("school", ""))
        body = {
            "staff_id": values.get("staff_id", ""),
            "name": values.get("name", ""),
            "email": values.get("email", ""),
            "designation": next(iter(_ref_codes(db, "designation", values.get("designation", ""), "designation")), ""),
            "subjects": _ref_codes(db, "subject", values.get("subjects", ""), "subjects"),
            "classes": _ref_codes(db, "class", values.get("classes", ""), "classes"),
            "experience_years": values.get("experience_years") or None,
            "qualifications": _qualifications(db, values.get("qualifications", "")),
        }
        user = _user_by_staff_id(db, body["staff_id"], Role.TEACHER)
        if user is None:
            people.create_teacher(
                db,
                principal,
                TeacherIn(school_id=school.id, custom=values["custom"], locale=values.get("locale") or None, **body),
                send_at=send_at,
                commit=False,
            )
            return "created"
        teacher = db.execute(select(Teacher).where(Teacher.user_id == user.id)).scalar_one()
        current_school = db.get(School, teacher.school_id)
        allowed = managed_cluster_ids(db, principal)
        if allowed is not None and current_school.cluster_id not in allowed:
            raise RowError([{"field": "staff_id", "code": "out_of_scope"}])
        if teacher.school_id != school.id:
            # Imports never move teachers; transfers are recorded explicitly (FR-T3).
            raise RowError([{"field": "school", "code": "use_transfer"}])
        before = people.teacher_out(db, teacher)
        email = body.pop("email")
        body.pop("staff_id")
        experience = body.pop("experience_years")
        update = TeacherUpdate(
            version=teacher.version,
            email=email if email != user.email else None,
            experience_years=experience
            if experience is not None and int(experience) != teacher.experience_years
            else None,
            custom=values["custom"] or None,
            **body,
        )
        people.update_teacher(db, principal, teacher.id, update, commit=False)
        return _changed(before, people.teacher_out(db, teacher))

    cluster = _cluster_by_code(db, principal, values.get("cluster", ""))
    user = _user_by_staff_id(db, values.get("staff_id", ""), Role.SUPERVISOR)
    if user is None:
        people.create_supervisor(
            db,
            principal,
            SupervisorIn(
                staff_id=values.get("staff_id", ""),
                name=values.get("name", ""),
                email=values.get("email", ""),
                cluster_id=cluster.id,
                locale=values.get("locale") or None,
            ),
            send_at=send_at,
            commit=False,
        )
        return "created"
    supervisor = db.execute(select(Supervisor).where(Supervisor.user_id == user.id)).scalar_one()
    allowed = managed_cluster_ids(db, principal)
    if allowed is not None and supervisor.cluster_id not in allowed:
        raise RowError([{"field": "staff_id", "code": "out_of_scope"}])
    if supervisor.cluster_id != cluster.id:
        raise RowError([{"field": "cluster", "code": "use_cluster_move"}])
    before = people.supervisor_out(db, supervisor)
    email = values.get("email", "")
    people.update_supervisor(
        db,
        principal,
        supervisor.id,
        PersonUpdate(
            version=supervisor.version,
            name=values.get("name") or None,
            email=email if email and email.lower() != user.email else None,
        ),
        commit=False,
    )
    return _changed(before, people.supervisor_out(db, supervisor))


def _changed(before: dict[str, Any], after: dict[str, Any]) -> str:
    strip = ("version",)
    return (
        "unchanged"
        if {k: v for k, v in before.items() if k not in strip} == {k: v for k, v in after.items() if k not in strip}
        else "updated"
    )


def _errors_from(exc: Exception) -> list[dict[str, Any]]:
    if isinstance(exc, RowError):
        return exc.errors
    if isinstance(exc, ValidationError):
        return [{"field": ".".join(str(p) for p in e["loc"]), "code": e["type"]} for e in exc.errors()]
    if isinstance(exc, AppError):
        if isinstance(exc.detail, list):
            return [{"code": exc.code, **d} if isinstance(d, dict) else {"code": exc.code} for d in exc.detail]
        return [{"code": exc.code, **({"detail": exc.detail} if exc.detail else {})}]
    raise exc


def _try_row(db: Session, principal: Principal, entity: str, values: dict[str, Any], *, keep: bool, send_at=None):
    savepoint = db.begin_nested()
    try:
        action = _apply_row(db, principal, entity, values, send_at=send_at)
    except (RowError, ValidationError, AppError) as exc:
        savepoint.rollback()
        return "error", _errors_from(exc)
    if keep:
        savepoint.commit()
    else:
        savepoint.rollback()
    return action, []


def _ensure_importer(principal: Principal) -> None:
    if principal.role not in (Role.DISTRICT_ADMIN, Role.CLUSTER_HEAD):
        raise AppError(403, "forbidden")


def preview(db: Session, principal: Principal, entity: str, filename: str, content: bytes) -> ImportBatch:
    _ensure_importer(principal)
    if entity not in ENTITIES:
        raise not_found()
    rows = parse(db, entity, filename, content)
    results = []
    seen: Counter[str] = Counter()
    key_field = "code" if entity == "schools" else "staff_id"
    for row in rows:
        key = row["values"].get(key_field, "")
        seen[key] += 1
        if key and seen[key] > 1:
            action, errors = "error", [{"field": key_field, "code": "duplicate_in_file"}]
        else:
            action, errors = _try_row(db, principal, entity, row["values"], keep=False)
        results.append({**row, "key": key, "action": action, "errors": errors})
    summary = dict(Counter(r["action"] for r in results))
    batch = ImportBatch(
        entity=entity, created_by=principal.user.id, filename=filename[:255], rows=results, summary=summary
    )
    db.add(batch)
    db.commit()
    return batch


def commit(db: Session, principal: Principal, batch_id: int) -> ImportBatch:
    _ensure_importer(principal)
    batch = db.execute(select(ImportBatch).where(ImportBatch.id == batch_id).with_for_update()).scalar_one_or_none()
    if batch is None or batch.created_by != principal.user.id:
        raise not_found()
    if batch.status != "previewed":
        raise AppError(409, "import_already_committed")
    pending = [r for r in batch.rows if r["action"] in ("created", "updated")]
    creates = sum(1 for r in pending if r["action"] == "created")
    # FR-N8: invitations from a bulk import are spread over days within the budget.
    send_times = iter(mail.invitation_send_times(db, creates)) if entity_has_users(batch.entity) else iter(())
    results = []
    for row in batch.rows:
        if row["action"] not in ("created", "updated"):
            results.append(row)
            continue
        send_at = next(send_times, None) if row["action"] == "created" else None
        action, errors = _try_row(db, principal, batch.entity, row["values"], keep=True, send_at=send_at)
        results.append({**row, "action": action, "errors": errors, "send_at": send_at.isoformat() if send_at else None})
    batch.rows = results
    batch.summary = dict(Counter(r["action"] for r in results))
    batch.status = "committed"
    batch.committed_at = clock.now()
    audit.record(
        db,
        "import.committed",
        entity_type="import",
        entity_id=batch.id,
        after={
            "entity": batch.entity,
            "summary": batch.summary,
        },
    )
    db.commit()
    return batch


def entity_has_users(entity: str) -> bool:
    return entity in ("teachers", "supervisors")


def batch_out(db: Session, batch: ImportBatch) -> dict[str, Any]:
    district = get_district(db)
    by_day: Counter[str] = Counter()
    for row in batch.rows:
        if row.get("send_at"):
            by_day[clock.local_date(district.time_zone, datetime.fromisoformat(row["send_at"])).isoformat()] += 1
    return {
        "id": batch.id,
        "entity": batch.entity,
        "filename": batch.filename,
        "status": batch.status,
        "summary": batch.summary,
        "rows": batch.rows,
        "invitation_schedule": [{"date": day, "count": count} for day, count in sorted(by_day.items())],
    }


# Export ------------------------------------------------------------------------------------


def export_rows(db: Session, principal: Principal, entity: str, locale: str) -> tuple[list[str], list[list[Any]]]:
    _ensure_importer(principal)
    if entity not in ENTITIES:
        raise not_found()
    district = get_district(db)
    columns = COLUMNS[entity]
    custom = active_custom_fields(db, CUSTOM_ENTITY[entity]) if entity in CUSTOM_ENTITY else []
    label = (lambda c: c.ar) if locale == "ar" else (lambda c: c.en)
    header = ["District" if locale != "ar" else "المنطقة"] + [label(c) for c in columns]
    header += [(f.label_ar if locale == "ar" else f.label_en) for f in custom]
    name = "name_ar" if locale == "ar" else "name_en"
    rows: list[list[Any]] = []
    page = 1
    while True:
        if entity == "schools":
            chunk = service.list_schools(db, principal, page=page, page_size=200)
            for s in chunk["items"]:
                rows.append(
                    [
                        district.code,
                        s["cluster"]["code"],
                        s["code"],
                        s["name"],
                        s["address"],
                        s["school_type"],
                        s["medium"],
                        s["contact_phone"],
                        s["contact_email"] or "",
                    ]
                    + [_custom_cell(s["custom"].get(f.key)) for f in custom]
                )
        elif entity == "teachers":
            chunk = people.list_teachers(db, principal, page=page, page_size=200)
            for t in chunk["items"]:
                rows.append(
                    [
                        district.code,
                        t["staff_id"],
                        t["name"],
                        t["email"],
                        t["school"]["code"],
                        t["designation"][name],
                        "; ".join(s[name] for s in t["subjects"]),
                        "; ".join(c[name] for c in t["classes"]),
                        t["experience_years"],
                        "; ".join(f"{q[name]}:{q['year']}" if q["year"] else q[name] for q in t["qualifications"]),
                        "",
                    ]
                    + [_custom_cell(t["custom"].get(f.key)) for f in custom]
                )
        else:
            chunk = people.list_supervisors(db, principal, page=page, page_size=200)
            for s in chunk["items"]:
                rows.append([district.code, s["staff_id"], s["name"], s["email"], s["cluster"]["code"], ""])
        if page * 200 >= chunk["total"]:
            break
        page += 1
    audit.record(
        db, "export.created", entity_type="export", entity_id=entity, after={"rows": len(rows), "locale": locale}
    )
    db.commit()
    return header, rows


def _custom_cell(value: Any) -> Any:
    if isinstance(value, list):
        return "; ".join(str(v) for v in value)
    if isinstance(value, bool):
        return "yes" if value else "no"
    return "" if value is None else value


def to_csv(header: list[str], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows(rows)
    # A byte order mark makes Excel open Arabic text correctly.
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(header: list[str], rows: list[list[Any]], rtl: bool) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.sheet_view.rightToLeft = rtl
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def template(entity: str, locale: str) -> bytes:
    """An empty import file with the right headings (FR-I8)."""
    if entity not in ENTITIES:
        raise not_found()
    header = [c.ar if locale == "ar" else c.en for c in COLUMNS[entity]]
    return to_csv(header, [])
